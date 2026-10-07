#!/usr/bin/env python3
"""Export IQVIA source data using the column schemas in common_datamodel.xlsx."""

import argparse
import csv
from decimal import Decimal, InvalidOperation
from io import StringIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


OUTPUT_SEPARATOR = ";"
OUTPUT_ENCODING = "utf-8"
DATETIME_FORMAT = "%Y-%m-%d %H:%M:%S"

ROOT = Path(__file__).resolve().parent.parent
ENTITY_SHEETS = (
    "HOSPITAL",
    "EPISODE",
    "LAB",
    "TRANSFUSION",
    "HOSP_PHARMACY",
    "HOSP_AMB_PHARMACY",
    "AMB_PHARMACY",
    "VISCOELASTIC_TEST",
)

SOURCE_MAPPINGS = {
    "HOSPITAL": {
        "cnh_cd": "HOSPITAL",
        "cnh_nm": "HOSPITAL_DES",
    },
    "EPISODE": {
        "cnh_cd": "HOSPITAL",
        "sex_cd": "SEXO",
        "admission_dt": "FECING",
        "discharge_dt": "FECALT",
        "discharge_type_cd": "TIPALT",
        "start_intervention_dt": "FECINT1",
        **{f"d{number}": f"D{number}" for number in range(1, 21)},
        **{f"poad{number}": f"POA_{number:02d}" for number in range(1, 21)},
        **{f"proc{number}": f"P{number}" for number in range(1, 21)},
    },
}


def parse_arguments():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=ROOT / "IQVIA 4T2024 H.csv",
        help="Ruta del CSV IQVIA (por defecto: archivo de la raíz del repositorio).",
    )
    parser.add_argument(
        "--model",
        type=Path,
        default=ROOT / "common_datamodel.xlsx",
        help="Ruta de la plantilla CDM (por defecto: archivo de la raíz del repositorio).",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=ROOT / "output",
        help="Carpeta donde guardar los CSV (por defecto: output/).",
    )
    return parser.parse_args()


def read_source(path):
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-8", "latin1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError(f"No se pudo detectar la codificación de {path}")

    sample_lines = [line for line in text.splitlines() if line.strip()][:10]
    delimiter_scores = []
    for delimiter in ("|", ";", ",", "\t"):
        lengths = [
            len(next(csv.reader([line], delimiter=delimiter)))
            for line in sample_lines
        ]
        if lengths and lengths[0] > 1:
            consistent = sum(length == lengths[0] for length in lengths) / len(lengths)
            delimiter_scores.append((lengths[0], consistent, delimiter))
    if not delimiter_scores:
        raise ValueError(f"No se pudo detectar el separador de {path}")
    delimiter = max(delimiter_scores)[2]

    frame = pd.read_csv(
        StringIO(text),
        sep=delimiter,
        dtype=str,
        keep_default_na=False,
        na_filter=False,
    )
    return frame


def read_schemas(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets_by_name = {sheet.title.casefold(): sheet for sheet in workbook.worksheets}
    schemas = {}
    try:
        for entity in ENTITY_SHEETS:
            sheet = sheets_by_name.get(entity.casefold())
            if sheet is None:
                raise ValueError(
                    f"No se encontró la hoja {entity!r} en la plantilla {path}"
                )
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            header_indexes = {name: index for index, name in enumerate(header)}
            if "label" not in header_indexes:
                raise ValueError(f"La hoja {sheet.title!r} no contiene la columna 'label'")

            labels = []
            formats = {}
            for row in rows:
                label = row[header_indexes["label"]]
                if label is None or str(label) == "":
                    continue
                label = str(label)
                if label in formats:
                    raise ValueError(
                        f"Etiqueta CDM duplicada {label!r} en la hoja {sheet.title!r}"
                    )
                labels.append(label)
                format_index = header_indexes.get("format")
                formats[label] = (
                    row[format_index]
                    if format_index is not None and format_index < len(row)
                    else None
                )
            schemas[entity] = (labels, formats)
    finally:
        workbook.close()
    return schemas


def parse_datetime(values):
    parsed = pd.to_datetime(values, format="mixed", dayfirst=True, errors="coerce")
    return parsed.dt.strftime(DATETIME_FORMAT).fillna("")


def format_numeric(value):
    value = str(value).strip()
    if not value:
        return ""
    if "," in value and "." in value:
        if value.rfind(",") > value.rfind("."):
            value = value.replace(".", "").replace(",", ".")
        else:
            value = value.replace(",", "")
    elif "," in value:
        value = value.replace(",", ".")
    try:
        number = Decimal(value)
    except InvalidOperation:
        return value
    if not number.is_finite():
        return ""
    return format(number, "f")


def calculate_age(source):
    birth_date = pd.to_datetime(
        source["FECNAC"], format="mixed", dayfirst=True, errors="coerce"
    )
    admission_date = pd.to_datetime(
        source["FECING"], format="mixed", dayfirst=True, errors="coerce"
    )
    age = admission_date.dt.year - birth_date.dt.year
    birthday_not_reached = (admission_date.dt.month < birth_date.dt.month) | (
        (admission_date.dt.month == birth_date.dt.month)
        & (admission_date.dt.day < birth_date.dt.day)
    )
    age = age - birthday_not_reached.astype("int64")
    return age.where(birth_date.notna() & admission_date.notna(), "").astype(str).replace(
        "<NA>", ""
    )


def build_entity(source, entity, schema):
    labels, formats = schema
    mappings = SOURCE_MAPPINGS.get(entity, {})
    available_mappings = {
        label: source_name
        for label, source_name in mappings.items()
        if label in labels and source_name in source.columns
    }
    has_derived_age = entity == "EPISODE" and "age_nm" in labels and {
        "FECNAC",
        "FECING",
    }.issubset(source.columns)
    if not available_mappings and not has_derived_age:
        return pd.DataFrame(columns=labels)

    result = pd.DataFrame(index=source.index)
    if has_derived_age:
        result["age_nm"] = calculate_age(source)
    for label, source_name in available_mappings.items():
        values = source[source_name].astype(str)
        field_format = str(formats.get(label) or "").casefold()
        if field_format in {"date", "datetime"}:
            result[label] = parse_datetime(values)
        elif field_format in {"integer", "double", "float", "decimal"}:
            result[label] = values.map(format_numeric)
        else:
            result[label] = values

    for label in labels:
        if label not in result:
            result[label] = ""
    result = result.loc[:, labels]
    if entity == "HOSPITAL":
        result = result.drop_duplicates(subset=["cnh_cd"], keep="first")
    return result


def generate(source_path, model_path, output_dir):
    source = read_source(source_path)
    schemas = read_schemas(model_path)
    output_dir.mkdir(parents=True, exist_ok=True)

    for entity in ENTITY_SHEETS:
        frame = build_entity(source, entity, schemas[entity])
        output_path = output_dir / f"{entity}.csv"
        frame.to_csv(
            output_path,
            sep=OUTPUT_SEPARATOR,
            encoding=OUTPUT_ENCODING,
            index=False,
        )
        if output_path.read_bytes().startswith(b"\xef\xbb\xbf"):
            raise AssertionError(f"Se generó un BOM no permitido en {output_path}")
        print(f"{output_path}: {len(frame)} filas, {len(frame.columns)} columnas")


def main():
    arguments = parse_arguments()
    generate(arguments.source, arguments.model, arguments.output_dir)


if __name__ == "__main__":
    main()
